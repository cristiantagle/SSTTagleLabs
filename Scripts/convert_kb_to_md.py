import datetime
import os
import pathlib
import sys
from typing import List

try:
    import fitz  # PyMuPDF
    import docx  # python-docx
    import pandas as pd  # type: ignore
    from pptx import Presentation  # python-pptx
except ImportError as exc:
    sys.stderr.write(f"Missing dependencies: {exc}. Run this script after installing requirements.\n")
    sys.exit(1)


SOURCE_ROOT = pathlib.Path("KnowledgeInput")
TARGET_ROOT = pathlib.Path("KnowledgeInput_md")

# Limits to avoid gigantic markdown files
MAX_XLSX_ROWS = 800  # per sheet
MAX_PDF_PAGES = 500


def ensure_parent(path: pathlib.Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def read_pdf(path: pathlib.Path) -> str:
    text_parts: List[str] = []
    with fitz.open(path) as doc:
        for page_idx, page in enumerate(doc):
            if page_idx >= MAX_PDF_PAGES:
                text_parts.append(f"\n\n[Truncated after {MAX_PDF_PAGES} pages]\n")
                break
            text_parts.append(f"# Page {page_idx + 1}\n")
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts)


def read_docx(path: pathlib.Path) -> str:
    document = docx.Document(path)
    parts: List[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Tables as markdown
    for table_idx, table in enumerate(document.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.replace("\n", " ").strip() for cell in row.cells])
        if rows:
            header = "| " + " | ".join(rows[0]) + " |"
            separator = "| " + " | ".join("---" for _ in rows[0]) + " |"
            body_rows = ["| " + " | ".join(r) + " |" for r in rows[1:]]
            table_md = "\n".join([f"## Table {table_idx + 1}", header, separator, *body_rows])
            parts.append(table_md)
    return "\n\n".join(parts)


def read_xlsx(path: pathlib.Path) -> str:
    parts: List[str] = []
    xls = pd.ExcelFile(path)
    for sheet_name in xls.sheet_names:
        df = xls.parse(sheet_name)
        if len(df) > MAX_XLSX_ROWS:
            df = df.head(MAX_XLSX_ROWS)
            note = f"_Truncated to {MAX_XLSX_ROWS} rows_"
        else:
            note = ""
        parts.append(f"# Sheet: {sheet_name}\n")
        parts.append(df.to_markdown(index=False))
        if note:
            parts.append(note)
    return "\n\n".join(parts)


def read_pptx(path: pathlib.Path) -> str:
    prs = Presentation(path)
    parts: List[str] = []
    for idx, slide in enumerate(prs.slides):
        slide_lines: List[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_lines.append(shape.text)
        if slide_lines:
            parts.append(f"# Slide {idx + 1}\n- " + "\n- ".join(slide_lines))
    return "\n\n".join(parts)


def convert_file(src: pathlib.Path) -> None:
    rel = src.relative_to(SOURCE_ROOT)
    target = TARGET_ROOT / rel
    target = target.with_suffix(".md")
    ensure_parent(target)

    ext = src.suffix.lower()
    content = ""
    kind = ext.upper().lstrip(".")

    try:
        if ext == ".pdf":
            content = read_pdf(src)
        elif ext in [".docx", ".doc"]:
            content = read_docx(src)
        elif ext in [".xlsx", ".xls"]:
            content = read_xlsx(src)
        elif ext in [".pptx", ".ppt"]:
            content = read_pptx(src)
        else:
            # Skip unsupported types (images, etc.) but create a stub
            content = f"(Archivo {ext} no convertido automáticamente. Ruta: {src})"
            kind = "BINARY"
    except Exception as exc:
        content = f"(ERROR al convertir {src}: {exc})"
        kind = "ERROR"

    meta = [
        f"# Fuente: {rel}",
        f"- Tipo: {kind}",
        f"- Última modificación: {datetime.datetime.fromtimestamp(src.stat().st_mtime)}",
        f"- Tamaño: {src.stat().st_size} bytes",
        "",
    ]
    target.write_text("\n".join(meta) + content, encoding="utf-8")


def main() -> None:
    if not SOURCE_ROOT.exists():
        sys.stderr.write(f"No se encontró {SOURCE_ROOT}\n")
        sys.exit(1)

    for src in SOURCE_ROOT.rglob("*"):
        if src.is_dir():
            continue
        if src.suffix.lower() in [".pdf", ".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt"]:
            convert_file(src)
        else:
            # Saltar imágenes y otros binarios; solo registramos stub
            convert_file(src)


if __name__ == "__main__":
    main()
